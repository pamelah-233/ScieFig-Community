"""Excel export for the application's measurement result tables."""
from math import isfinite
from pathlib import Path

from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QFileDialog, QMessageBox, QPushButton, QTableWidget

from core.utils.i18n import bind_i18n, tr, trf
from core.utils.recent_directory import recent_directory, remember_file_directory

EXCEL_VALUE_ROLE = int(Qt.ItemDataRole.UserRole) + 1


def _write_table(sheet, table: QTableWidget, numeric_columns: tuple[int, ...]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    for col in range(table.columnCount()):
        header = table.horizontalHeaderItem(col)
        cell = sheet.cell(1, col + 1, header.text() if header else str(col + 1))
        cell.data_type = "s"
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="304B62")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in range(table.rowCount()):
        for col in range(table.columnCount()):
            item = table.item(row, col)
            text = item.text() if item else ""
            value = text or None
            if col in numeric_columns and item is not None:
                raw = item.data(EXCEL_VALUE_ROLE)
                try:
                    number = float(raw if raw is not None else text)
                except (ValueError, TypeError):
                    pass
                else:
                    if isfinite(number):
                        value = int(number) if number.is_integer() else number
            cell = sheet.cell(row + 2, col + 1, value)
            # A sample/image name beginning with '=' is text, never a formula.
            if isinstance(value, str):
                cell.data_type = "s"
            elif isinstance(value, float):
                cell.number_format = "0.############"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    for col in sheet.columns:
        width = max(sum(2 if ord(ch) > 255 else 1 for ch in str(cell.value or "")) for cell in col)
        sheet.column_dimensions[get_column_letter(col[0].column)].width = min(48, max(14, width + 2))


def export_tables_xlsx(tables, path: str | Path, *, provenance=None) -> Path:
    from openpyxl import Workbook
    book = Workbook()
    book.remove(book.active)
    for title, table, numeric_columns in tables:
        if table.rowCount():
            _write_table(book.create_sheet(title), table, numeric_columns)
    if not book.worksheets:
        raise ValueError(tr("no_measurements_to_export"))
    if provenance:
        import json
        columns = list(dict.fromkeys(key for row in provenance for key in row))
        sheet = book.create_sheet('Provenance')
        sheet.append(columns)
        long_values = None
        for row in provenance:
            values = [json.dumps(row.get(key), ensure_ascii=False, default=str) if isinstance(row.get(key), (dict,list)) else row.get(key) for key in columns]
            for column, value in enumerate(values):
                if isinstance(value, str) and len(value) > 32767:
                    if long_values is None:
                        long_values = book.create_sheet('Provenance values')
                        long_values.append(['Provenance row', 'Column', 'Part', 'Text'])
                    for part, start in enumerate(range(0, len(value), 30000), 1):
                        long_values.append([sheet.max_row + 1, columns[column], part, value[start:start+30000]])
                        long_values.cell(long_values.max_row, 4).data_type = 's'
                    values[column] = f"[Provenance values: row {sheet.max_row + 1}, {columns[column]}]"
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    target = Path(path)
    book.save(target)
    book.close()
    return target


def export_table_xlsx(table: QTableWidget, path: str | Path, numeric_columns: tuple[int, ...]) -> Path:
    return export_tables_xlsx([("Measurements", table, numeric_columns)], path)


def measurement_excel_button(owner, tables, filename: str) -> QPushButton:
    button = bind_i18n(QPushButton(), "export_excel")
    bind_i18n(button, "export_measurement_table_tooltip", "toolTip")

    def export() -> None:
        parent = owner.window()
        if not any(table.rowCount() for _, table, _ in tables):
            QMessageBox.information(parent, tr("export_excel"), tr("no_measurements_to_export"))
            return
        selected, _ = QFileDialog.getSaveFileName(
            parent, tr("export_excel"), str(Path(recent_directory()) / filename), "Excel (*.xlsx)"
        )
        if not selected:
            return
        target = Path(selected)
        if target.suffix.lower() != ".xlsx":
            target = target.with_suffix(".xlsx")
        try:
            provider = getattr(button, 'provenance_provider', None)
            export_tables_xlsx(tables, target, provenance=provider() if provider else None)
        except (OSError, ValueError) as error:
            QMessageBox.warning(parent, tr("export_excel"), trf("measurement_excel_failed", error=str(error)))
            return
        remember_file_directory(target)
        QMessageBox.information(parent, tr("export_excel"), trf("measurement_excel_saved", path=str(target)))

    button.clicked.connect(export)
    return button


def table_excel_button(table: QTableWidget, numeric_columns: tuple[int, ...], filename: str) -> QPushButton:
    return measurement_excel_button(table, [("Measurements", table, numeric_columns)], filename)
