from threading import Event, get_ident
from time import monotonic

from openpyxl import load_workbook
from PyQt6.QtCore import QSettings, Qt
from PyQt6.QtWidgets import QApplication, QFileDialog, QMessageBox
from PyQt6.QtTest import QTest

from core.utils.calibration_presets import CalibrationPresets
from tests.test_batch_analysis import fluorescence_panel


def app_instance():
    return QApplication.instance() or QApplication([])


def wait_for(app, predicate):
    deadline = monotonic() + 10
    while not predicate() and monotonic() < deadline:
        app.processEvents()
    assert predicate(), "batch UI did not finish"


def test_dialog_checks_selection_runs_in_background_and_refreshes_on_gui_thread(tmp_path):
    from ui.dialogs.batch_analysis_dialog import BatchAnalysisDialog

    app = app_instance()
    panel = fluorescence_panel()
    main_thread = get_ident()
    analysis_threads, refresh_threads = [], []
    library = CalibrationPresets(QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat))
    library.save("Camera preset", {"um_per_px_x": .5, "um_per_px_y": .7})

    def runner(snapshot, uid, method, options):
        analysis_threads.append(get_ident())
        return getattr(snapshot, method)(uid, **options)

    dialog = BatchAnalysisDialog(panel, "count_channel_objects", {"channel_indices": [0]},
        on_refresh=lambda: refresh_threads.append(get_ident()), runner=runner, library=library)
    dialog.image_table.item(2, 0).setCheckState(Qt.CheckState.Unchecked)
    dialog.preset_combo.setCurrentIndex(1)
    dialog.start_button.click()
    wait_for(app, lambda: not dialog.is_running)

    assert len(analysis_threads) == 2
    assert all(thread != main_thread for thread in analysis_threads)
    assert len(refresh_threads) == 3  # Each success, plus completion.
    assert set(refresh_threads) == {main_thread}
    assert len(panel.analysis["object_counts"]) == 2
    assert panel.data["active_image_uid"] == "image-2"
    assert panel.measurement_scale_bar("image-0")["um_per_px_y"] == .7
    assert dialog.progress_bar.value() == 2
    assert dialog.export_button.isEnabled()
    dialog.close()
    panel.dispose()


def test_dialog_cancel_waits_for_current_image_and_preserves_success(tmp_path):
    from ui.dialogs.batch_analysis_dialog import BatchAnalysisDialog

    app = app_instance()
    panel = fluorescence_panel()
    entered, release = Event(), Event()
    library = CalibrationPresets(QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat))

    def runner(snapshot, uid, method, options):
        entered.set()
        assert release.wait(5)
        return getattr(snapshot, method)(uid, **options)

    dialog = BatchAnalysisDialog(panel, "count_channel_objects", {"channel_indices": [0]}, runner=runner, library=library)
    dialog.show()
    dialog.start_button.click()
    wait_for(app, entered.is_set)
    running_status = dialog.image_table.item(0, 2).text()
    dialog.cancel_button.click()
    dialog.reject()
    assert dialog.isVisible()
    release.set()
    wait_for(app, lambda: not dialog.is_running)
    assert running_status in {"Running", "运行中"}
    assert [result.status for result in dialog.results] == ["success", "cancelled", "cancelled"]
    assert len(panel.analysis["object_counts"]) == 1
    assert dialog.export_button.isEnabled()
    dialog.close()
    panel.dispose()


def test_dialog_empty_selection_and_export_cancellation_write_nothing(tmp_path, monkeypatch):
    from ui.dialogs.batch_analysis_dialog import BatchAnalysisDialog

    app = app_instance()
    panel = fluorescence_panel()
    library = CalibrationPresets(QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat))
    dialog = BatchAnalysisDialog(panel, "count_channel_objects", {"channel_indices": [0]}, library=library)
    dialog.select_none_button.click()
    assert not dialog.start_button.isEnabled()
    dialog.select_all_button.click()
    dialog.start_button.click()
    wait_for(app, lambda: not dialog.is_running)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: ("", ""))
    dialog.export_button.click()
    assert not list(tmp_path.glob("*.xlsx"))
    dialog.close()
    panel.dispose()


def test_export_contains_numeric_results_settings_calibration_ids_and_current_review(tmp_path):
    from core.batch_analysis import BatchAnalysisJob, apply_batch_image_result
    from ui.dialogs.batch_analysis_dialog import export_batch_results_xlsx

    app = app_instance()
    panel = fluorescence_panel()
    panel.images()[0]["source_name"] = "=SUM(1,2).tif"
    results = BatchAnalysisJob(panel, ["image-0"], "count_channel_objects", {"channel_indices": [0]}).run()
    apply_batch_image_result(panel, results[0])
    live = panel.analysis["object_counts"][0]
    live["review_status"] = "edited"
    live["count"] = 1
    live["object_review"] = {"version": 1, "audit": [{"action": "delete", "at": "2026-09-08T00:00:00Z"}]}
    live["detected_spots"][0]["review_provenance"] = {"status": "edited", "action": "replace"}
    path = export_batch_results_xlsx(panel, results, tmp_path / "batch.xlsx")
    workbook = load_workbook(path, data_only=False)
    assert {"Images", "Measurements", "Objects"} <= set(workbook.sheetnames)
    sheet = workbook["Measurements"]
    row = dict(zip([cell.value for cell in sheet[1]], [cell.value for cell in sheet[2]]))
    assert row["image_uid"] == "image-0"
    assert row["measurement_id"] == live["count_uuid"]
    assert row["method"] == "count_channel_objects"
    assert row["count"] == 1
    assert isinstance(row["count"], int)
    assert '"channel_indices": [0]' in row["settings"]
    assert '"um_per_px_y": 0.6' in row["calibration"]
    assert row["source_hash"] == "stored-source-0"
    assert row["review_status"] == "edited"
    assert '"delete"' in row["object_review"]
    source_cell = sheet.cell(2, [cell.value for cell in sheet[1]].index("source_name") + 1)
    assert source_cell.value == "=SUM(1,2).tif"
    assert source_cell.data_type == "s"
    objects = workbook["Objects"]
    object_row = dict(zip([cell.value for cell in objects[1]], [cell.value for cell in objects[2]]))
    assert '"replace"' in object_row["review_provenance"]
    workbook.close()
    panel.dispose()


def test_export_long_review_provenance_is_lossless_in_excel(tmp_path):
    import json
    from core.batch_analysis import BatchAnalysisJob, apply_batch_image_result
    from ui.dialogs.batch_analysis_dialog import export_batch_results_xlsx

    app = app_instance()
    panel = fluorescence_panel()
    results = BatchAnalysisJob(panel, ["image-0"], "count_channel_objects", {"channel_indices": [0]}).run()
    apply_batch_image_result(panel, results[0])
    review = {"audit": [{"action": "merge", "after": "review contour " * 3000}]}
    panel.analysis["object_counts"][0]["object_review"] = review
    path = export_batch_results_xlsx(panel, results, tmp_path / "long.xlsx")
    workbook = load_workbook(path)
    parts = [dict(zip([c.value for c in workbook["Long values"][1]], row))
             for row in workbook["Long values"].iter_rows(min_row=2, values_only=True)]
    text = "".join(part["value"] for part in sorted(parts, key=lambda part: part["part"])
                   if part["field"] == "object_review")
    assert json.loads(text) == review
    workbook.close()
    panel.dispose()


def test_full_image_scope_is_explicit_and_image_checkboxes_work_with_keyboard(tmp_path):
    from ui.dialogs.batch_analysis_dialog import BatchAnalysisDialog

    app = app_instance()
    panel = fluorescence_panel()
    library = CalibrationPresets(QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat))
    dialog = BatchAnalysisDialog(panel, "count_channel_objects", {"channel_indices": [0]}, library=library)
    assert dialog.scope_label.text() in {"Scope: full image", "范围：整幅图像"}
    dialog.show()
    dialog.image_table.setCurrentCell(0, 0)
    dialog.image_table.setFocus()
    QTest.keyClick(dialog.image_table, Qt.Key.Key_Space)
    assert len(dialog.selected_image_uids()) == 2
    dialog.close()
    panel.dispose()


def test_export_suffix_normalization_does_not_overwrite_existing_target_without_confirmation(tmp_path, monkeypatch):
    from ui.dialogs.batch_analysis_dialog import BatchAnalysisDialog

    app = app_instance()
    panel = fluorescence_panel()
    library = CalibrationPresets(QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat))
    dialog = BatchAnalysisDialog(panel, "count_channel_objects", {"channel_indices": [0]}, library=library)
    dialog.start_button.click()
    wait_for(app, lambda: not dialog.is_running)
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"existing report")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: (str(target.with_suffix("")), ""))
    monkeypatch.setattr(QMessageBox, "question", lambda *args, **kwargs: QMessageBox.StandardButton.No)
    dialog.export_button.click()
    assert target.read_bytes() == b"existing report"
    dialog.close()
    panel.dispose()
