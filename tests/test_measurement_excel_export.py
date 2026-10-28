import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import load_workbook
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QApplication, QTableWidget, QTableWidgetItem

from ui.widgets import measurement_excel_export as excel


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


def test_table_export_preserves_numbers_units_and_literal_identifiers(app, tmp_path):
    table = QTableWidget(2, 4)
    table.setHorizontalHeaderLabels(["图像", "面积 (µm²)", "直径 (µm)", "分组"])
    for row, values in enumerate([["001", "123.45", "0", "=1+1"], ["细胞.tif", "", "1.23e-4", "A"]]):
        for col, text in enumerate(values):
            table.setItem(row, col, QTableWidgetItem(text))
    table.item(0, 1).setData(excel.EXCEL_VALUE_ROLE, 123.456789)
    path = excel.export_table_xlsx(table, tmp_path / "结果.xlsx", (1, 2))
    book = load_workbook(path)
    ws = book.active
    assert ws["B1"].value == "面积 (µm²)"
    assert ws["A2"].value == "001"
    assert ws["B2"].value == 123.456789
    assert ws["C2"].value == 0
    assert ws["B3"].value is None
    assert ws["C3"].value == pytest.approx(0.000123)
    assert ws["D2"].data_type == "s"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:D3"
    book.close()


def test_button_cancel_and_empty_do_not_write(app, tmp_path, monkeypatch):
    table = QTableWidget(0, 1)
    table.setHorizontalHeaderLabels(["Count"])
    messages = []
    monkeypatch.setattr(excel.QMessageBox, "information", lambda *args: messages.append(args))
    monkeypatch.setattr(excel.QFileDialog, "getSaveFileName", lambda *args, **kw: ("", ""))
    button = excel.table_excel_button(table, (0,), "counts.xlsx")
    button.click()
    assert messages
    table.setRowCount(1)
    table.setItem(0, 0, QTableWidgetItem("5"))
    button.click()
    assert not list(tmp_path.iterdir())


def test_measurement_panels_have_working_excel_buttons(app, tmp_path, monkeypatch):
    from ui.panels.placeholder_panel import DataInputPanel
    from ui.panels.fluorescence_workspace import FluorescenceWorkspace
    data = DataInputPanel()
    fluoro = FluorescenceWorkspace()
    fluoro.results_drawer.setChecked(True)
    groups = [
        (data, "ihc", ["ihc_results"]),
        (data, "brightfield", ["brightfield_results", "brightfield_ai_results", "brightfield_ai_fiber_area", "brightfield_bright_round_cell_results"]),
        (fluoro, "fluoro", ["measurement", "count", "overlap", "ai_results", "ai_summary_stats"]),
    ]
    monkeypatch.setattr(excel.QMessageBox, "information", lambda *args: None)
    monkeypatch.setattr(excel, "remember_file_directory", lambda path: None)
    for owner, module, prefixes in groups:
        for prefix in prefixes:
            table = getattr(owner, prefix + "_table")
            table.setRowCount(1)
            for col in range(table.columnCount()):
                table.setItem(0, col, QTableWidgetItem("7.25"))
        path = tmp_path / (module + ".xlsx")
        monkeypatch.setattr(excel.QFileDialog, "getSaveFileName", lambda *args, p=path, **kw: (str(p), ""))
        getattr(owner, module + "_measurement_excel_button").click()
        book = load_workbook(path)
        assert len(book.worksheets) == len(prefixes)
        for sheet in book.worksheets:
            assert sheet.max_row == 2
            assert any(cell.data_type == "n" for cell in sheet[2])
        book.close()


def test_one_excel_entry_and_one_analysis_panel_at_a_time(app):
    from ui.panels.placeholder_panel import DataInputPanel
    from ui.panels.fluorescence_workspace import FluorescenceWorkspace
    from PyQt6.QtWidgets import QPushButton
    from core.utils.i18n import tr
    data = DataInputPanel()
    fluoro = FluorescenceWorkspace()
    data.set_module("brightfield")
    data.set_brightfield_step(2, emit_signal=False)
    assert data.brightfield_measure_mode_stack.count() == 3
    for index in range(3):
        data.brightfield_measure_type_combo.setCurrentIndex(index)
        assert data.brightfield_measure_mode_stack.currentIndex() == index
    for root in (data.brightfield_measure_group, data.ihc_results_group, fluoro.context_stack.widget(2)):
        buttons = [button for button in root.findChildren(QPushButton) if button.text() == tr("export_excel")]
        assert len(buttons) == 1


def test_export_uses_original_measurements_not_rounded_display(app, tmp_path):
    from ui.panels.placeholder_panel import DataInputPanel
    from ui.panels.fluorescence_workspace import FluorescenceWorkspace
    from core.models import FluorescencePanel
    data = DataInputPanel()
    data.populate_ihc_measurements([{"results": {"positive_area_percent": 12.3456789}}])
    data.populate_brightfield_measurements([{"results": {"mean_object_area_px": 12.3456789}}])
    fluoro = FluorescenceWorkspace()
    fluoro.panel = FluorescencePanel(panel_id="export-test")
    fluoro.panel.data["fluorescence_images"] = [{"uid":"one", "measurements":[{
        "type":"roi_intensity", "results_table":[{"mean":0.000123456789}]
    }]}]
    fluoro.panel.data["active_image_uid"] = "one"
    fluoro._populate_measurements()
    for index, (table, column, expected) in enumerate([
        (data.ihc_results_table, 2, 12.3456789),
        (data.brightfield_results_table, 6, 12.3456789),
        (fluoro.measurement_table, 4, 0.000123456789),
    ]):
        path = excel.export_table_xlsx(table, tmp_path / f"precision{index}.xlsx", (column,))
        book = load_workbook(path)
        assert book.active.cell(2, column+1).value == expected
        book.close()
