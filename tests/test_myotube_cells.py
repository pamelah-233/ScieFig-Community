"""Native-contour myotube metrics, independent of model detection quality."""
import cv2
import numpy as np
import pytest

from core.image_analysis.cellpose_cells import measurements_from_labels
from core.image_analysis.cell_calibration import calibrate_cell_analysis


def test_long_cell_reports_transverse_width_not_feret_length():
    labels=np.zeros((130,340),np.int32)
    cv2.rectangle(labels,(25,35),(285,55),1,-1)
    result=measurements_from_labels(np.zeros((*labels.shape,3),np.uint8),labels,.5,{})
    row=result['results'][0]
    assert row['max_feret_um']==pytest.approx(np.hypot(260,20)*.5)
    assert row['max_width_px']==pytest.approx(20,abs=1)
    assert row['max_width_um']==pytest.approx(10,abs=.5)
    assert row['area_um2']==pytest.approx(row['area_px']*.25)
    assert row['max_feret_um']>3*row['equivalent_diameter_um']
    assert result['summary']['complete_count']==1
    assert result['summary']['mean_max_width_um']==row['max_width_um']
    assert result['summary']['mean_area_um2']==row['area_um2']


def test_truncated_myotube_does_not_bias_complete_size_summary():
    labels=np.zeros((140,340),np.int32)
    cv2.rectangle(labels,(25,35),(125,55),1,-1)
    cv2.rectangle(labels,(0,90),(300,115),2,-1)
    result=measurements_from_labels(np.zeros((*labels.shape,3),np.uint8),labels,.5,{})
    complete,truncated=result['results']
    assert truncated['border_truncated'] and not truncated['include_in_statistics']
    assert result['summary']['complete_count']==1
    assert result['summary']['mean_max_width_um']==complete['max_width_um']
    assert result['summary']['mean_area_um2']==complete['area_um2']


def test_myotube_area_and_length_honor_saved_anisotropic_calibration():
    labels=np.zeros((130,340),np.int32)
    cv2.rectangle(labels,(25,35),(285,55),1,-1)
    initial=measurements_from_labels(np.zeros((*labels.shape,3),np.uint8),labels,1.,{})
    result=calibrate_cell_analysis(initial,dict(um_per_px_x=.5,um_per_px_y=.25,calibrated=True))
    row=result['results'][0]
    assert row['area_um2']==pytest.approx(row['area_px']*.5*.25)
    assert row['max_feret_um']==pytest.approx(np.hypot(260*.5,20*.25))
    assert row['max_width_um']==pytest.approx(5,abs=.5)
    assert result['summary']['mean_max_width_um']==row['max_width_um']
    assert initial['results'][0]['max_feret_um']>row['max_feret_um']


def test_curved_myotube_width_does_not_use_its_bounding_box():
    from core.image_analysis import cell_width
    contour=np.zeros((230,230),np.uint8)
    points=np.array([[115+80*np.cos(t),115+80*np.sin(t)] for t in np.linspace(0,np.pi*1.3,180)],np.int32)
    cv2.polylines(contour,[points],False,1,16)
    result=cell_width.mask_width_metrics(contour,.5,.5)
    assert 7<result['max_width_um']<11
    assert np.linalg.norm(np.subtract(result['width_end_px'],result['width_start_px']))==pytest.approx(result['max_width_px'])


def test_local_swelling_sets_maximum_width_not_whole_cell_length():
    from core.image_analysis import cell_width
    mask=np.zeros((120,330),np.uint8)
    cv2.rectangle(mask,(20,48),(300,62),1,-1)
    cv2.circle(mask,(170,55),18,1,-1)
    result=cell_width.mask_width_metrics(mask,1.,1.)
    assert result['max_width_px']==pytest.approx(36,abs=2)


def test_disconnected_review_parts_do_not_measure_the_gap_as_cell_width():
    from core.image_analysis import cell_width
    mask=np.zeros((150,260),np.uint8)
    cv2.rectangle(mask,(20,20),(200,35),1,-1)
    cv2.rectangle(mask,(20,110),(200,125),1,-1)
    result=cell_width.mask_width_metrics(mask,1.,1.)
    assert result['max_width_px']==pytest.approx(15,abs=1)


def test_manual_contour_correction_recomputes_width_and_area_and_undo():
    from core.measurement_review import ObjectReviewSession
    labels=np.zeros((130,340),np.int32)
    cv2.rectangle(labels,(25,35),(285,55),1,-1)
    initial=measurements_from_labels(np.zeros((*labels.shape,3),np.uint8),labels,.5,{})
    record=dict(type='brightfield_bright_round_cell_diameter',results=initial['results'],summary=initial['summary'],
                measurement_scale_bar=dict(um_per_px=.5,um_per_px_x=.5,um_per_px_y=.5,calibrated=True))
    session=ObjectReviewSession(record,labels.shape)
    session.replace_polygon(1,[[25,35],[285,35],[285,75],[25,75]])
    edited=session.updated_record()
    assert edited['results'][0]['max_width_um']==pytest.approx(20,abs=1)
    assert edited['summary']['mean_max_width_um']==edited['results'][0]['max_width_um']
    assert edited['results'][0]['area_um2']>record['results'][0]['area_um2']
    assert session.undo()
    assert session.updated_record()['results'][0]['max_width_um']==record['results'][0]['max_width_um']


def test_cell_result_table_exposes_saved_transverse_width_and_area(qt_application,tmp_path,monkeypatch):
    from ui.panels.placeholder_panel import DataInputPanel
    from core.utils.i18n import tr
    from ui.widgets.measurement_excel_export import EXCEL_VALUE_ROLE
    panel=DataInputPanel()
    labels=np.zeros((130,340),np.int32);cv2.rectangle(labels,(25,35),(285,55),1,-1)
    result=measurements_from_labels(np.zeros((*labels.shape,3),np.uint8),labels,.5,{})
    panel.populate_bright_round_cell_results(dict(image_name='myotube.png',**result))
    table=panel.brightfield_bright_round_cell_results_table
    assert table.columnCount()==12
    assert table.horizontalHeaderItem(10).text()==tr('cell_max_transverse_width')
    assert table.item(0,10).data(EXCEL_VALUE_ROLE)==result['results'][0]['max_width_um']
    assert table.item(0,11).data(EXCEL_VALUE_ROLE)==result['results'][0]['area_um2']
    from ui.widgets import measurement_excel_export as excel
    from openpyxl import load_workbook
    monkeypatch.setattr(excel.QFileDialog,'getSaveFileName',lambda *args: (str(tmp_path/'width.xlsx'),''))
    monkeypatch.setattr(excel,'remember_file_directory',lambda *args:None)
    monkeypatch.setattr(excel.QMessageBox,'information',lambda *args:None)
    panel.brightfield_measurement_excel_button.click()
    book=load_workbook(tmp_path/'width.xlsx')
    sheet=book['Cell sizes']
    assert sheet.cell(2,11).value==pytest.approx(result['results'][0]['max_width_um'])
    assert sheet.cell(2,12).value==pytest.approx(result['results'][0]['area_um2'])
    book.close()
    panel.close()
