"""Independent myotube workflow; reuse transport and geometry, isolate records."""
from copy import deepcopy
import cv2
import numpy as np
import pytest


def example():
    labels=np.zeros((100,270),np.int32)
    cv2.rectangle(labels,(30,25),(220,45),1,-1)
    cv2.circle(labels,(50,75),10,2,-1)
    source=np.full(labels.shape,500,np.uint16);source[labels>0]=25000
    return source,labels


def test_myotube_native_channel_width_area_filter_and_calibration():
    from core.image_analysis.myotube_cells import myotube_measurements_from_labels
    source,labels=example();before=source.copy()
    result=myotube_measurements_from_labels(source,labels,modality='fluoro',settings={},
        calibration=dict(um_per_px_x=.5,um_per_px_y=.25,calibrated=True))
    assert result['type']=='myotube_measurement' and result['modality']=='fluoro'
    assert len(result['results'])==1  # The round nucleus-like object is not a tube.
    row=result['results'][0]
    assert row['max_width_um']==pytest.approx(5,abs=.5)
    assert row['area_um2']==pytest.approx(row['area_px']*.5*.25)
    assert row['max_width_um']<row['max_feret_um']/4
    assert row['mean_intensity']==25000
    assert row['centerline_length_px']==pytest.approx(190,abs=2)
    assert row['centerline_length_um']==pytest.approx(95,abs=1)
    assert result['summary']['mean_centerline_length_um']==row['centerline_length_um']
    assert result['summary']['complete_count']==1
    np.testing.assert_array_equal(source,before)


def test_myotube_inference_keeps_uint16_channel_and_effective_settings(monkeypatch):
    from core.image_analysis import myotube_cells,cellpose_backend
    source,labels=example();calls=[]
    def worker(image,**settings):
        calls.append((image.copy(),settings));return labels,dict(model='cyto3',diameter_px=settings['diameter'])
    monkeypatch.setattr(cellpose_backend,'run_cellpose_labels',worker)
    result=myotube_cells.analyze_myotubes(source,modality='fluoro',settings=dict(diameter_px=60,min_area_px=50),calibration={})
    np.testing.assert_array_equal(calls[0][0],source)
    assert calls[0][0].dtype==np.uint16 and calls[0][1]['diameter']==60
    assert result['provenance']['model']=='cyto3'
    assert result['settings']['diameter_px']==60
    row=result['results'][0]
    assert row['max_width_px']==pytest.approx(20,abs=1)
    assert row['max_width_um'] is None and row['area_um2'] is None
    assert row['centerline_length_um'] is None
    assert row['centerline_length_px']==pytest.approx(190,abs=2)


def test_myotube_truncated_objects_excluded_and_invalid_roi_rejected():
    from core.image_analysis.myotube_cells import myotube_measurements_from_labels
    source,labels=example();labels[0:10,0:240]=3
    result=myotube_measurements_from_labels(source,labels,modality='fluoro',settings={},calibration={})
    assert result['summary']['complete_count']==1
    assert result['summary']['border_truncated_count']==1
    assert not next(r for r in result['results'] if r['border_truncated'])['include_in_statistics']
    with pytest.raises(ValueError,match='ROI'):
        myotube_measurements_from_labels(source,labels,modality='fluoro',settings={},calibration={},roi_mask=np.ones((2,3)))


def test_myotube_manual_add_reuses_width_metrics():
    from core.measurement_review import ObjectReviewSession
    session=ObjectReviewSession(dict(type='myotube_measurement',results=[],summary={},
        measurement_scale_bar=dict(um_per_px=.5,calibrated=True)),(100,270))
    session.add_polygon([[20,20],[230,20],[230,40],[20,40]])
    row=session.updated_record()['results'][0]
    assert row['max_width_um']==pytest.approx(10,abs=1)
    assert row['centerline_length_um']==pytest.approx(105,abs=1)
    assert session.undo() and not session.updated_record()['results']


def test_broad_branch_is_flagged_not_pooled_as_one_complete_fiber():
    from core.image_analysis.myotube_cells import myotube_measurements_from_labels
    labels=np.zeros((300,500),np.int32)
    cv2.rectangle(labels,(30,160),(470,190),1,-1)
    cv2.rectangle(labels,(230,35),(260,170),1,-1)
    result=myotube_measurements_from_labels(np.uint16(labels*25000),labels,modality='fluoro',
        calibration=dict(um_per_px=.5,calibrated=True))
    row=result['results'][0]
    assert row['geometry_review_required']
    assert row['geometry_review_reason']=='branched_or_joined_contour'
    assert not row['include_in_statistics']
    assert result['summary']['mean_centerline_length_um'] is None
    from core.measurement_review import ObjectReviewSession
    reviewed=ObjectReviewSession(result,labels.shape)
    reviewed.replace_polygon(1,[[30,160],[470,160],[470,190],[30,190]])
    corrected=reviewed.updated_record()['results'][0]
    assert not corrected['geometry_review_required']
    assert corrected['include_in_statistics']
    assert reviewed.updated_record()['summary']['geometry_review_count']==0


def test_myotube_pixel_overlay_labels_width_not_length(monkeypatch):
    from core.image_analysis.myotube_cells import myotube_measurements_from_labels
    from core.image_analysis.bright_round_cells import render_bright_round_cell_overlay
    source,labels=example();result=myotube_measurements_from_labels(source,labels,modality='fluoro',calibration={})
    texts=[];original=cv2.putText
    def put_text(image,text,*args,**kwargs):
        texts.append(text);return original(image,text,*args,**kwargs)
    monkeypatch.setattr(cv2,'putText',put_text)
    render_bright_round_cell_overlay(np.zeros((*labels.shape,3),np.uint8),result)
    assert any('Wmax 20.0 px' in text for text in texts)
    assert any('L ' in text for text in texts)


def test_unmeasurable_width_is_not_drawn_as_feret_length(monkeypatch):
    from core.image_analysis.bright_round_cells import render_bright_round_cell_overlay
    lines=[]
    monkeypatch.setattr(cv2,'line',lambda *args,**kwargs:lines.append(args))
    row=dict(id=1,shape_class='elongated',max_width_px=None,max_width_um=None,
             feret_start_px=[10,10],feret_end_px=[190,10],max_feret_um=90,min_feret_um=4)
    render_bright_round_cell_overlay(np.zeros((60,220,3),np.uint8),dict(results=[row]))
    assert not lines, 'A missing transverse width must not silently become a longitudinal line.'


def test_calibrated_result_with_uncertain_width_keeps_area_and_length_visible(qt_application):
    from ui.widgets.myotube_controls import MyotubeControls
    from core.utils.i18n import tr
    controls=MyotubeControls('fluoro')
    result=dict(calibrated=True,results=[dict(id=1,max_width_px=None,area_px=200.,area_um2=50.,
        centerline_length_px=100.,centerline_length_um=50.)],summary=dict(count=1,complete_count=1,
        mean_max_width_um=None,mean_area_um2=50.,mean_centerline_length_um=50.))
    controls.show_result(result)
    assert tr('myotube_pixel_hint') not in controls.summary.text()
    assert tr('myotube_width_review') in controls.table.item(0,7).text()
    assert controls.table.item(0,6).text()=='50.00'
    controls.close()


def test_myotube_controls_do_not_describe_signal_as_dapi_or_area_as_nuclei(qt_application):
    from ui.widgets.myotube_controls import MyotubeControls
    from PyQt6.QtWidgets import QLabel
    from core.utils.i18n import set_language,get_language
    previous=get_language()
    try:
        set_language('zh');controls=MyotubeControls('fluoro')
        labels=[label.text() for label in controls.findChildren(QLabel)]
        assert not any('DAPI' in text or '核面积' in text for text in labels)
        assert any('肌管信号通道' in text for text in labels)
        assert any('最小肌管面积' in text for text in labels)
        controls.close()
    finally:set_language(previous)


@pytest.fixture
def myotube_window(qt_application,desktop_screen):
    from pathlib import Path
    from ui.main_window import MainWindow
    window=MainWindow(Path(__file__).resolve().parents[1]/'resources')
    window.project_recovery.timer.stop()
    yield window
    window._capture_clean_project_baseline();window.close();qt_application.processEvents()


def import_example(window,module,tmp_path):
    import tifffile
    from PIL import Image
    controller=window.myotube_controllers[module];panel=controller.panel()
    source,labels=example()
    if module=='fluoro':
        path=tmp_path/'myhc.tif';tifffile.imwrite(path,np.stack([source,source*2]),metadata={'axes':'CYX'},photometric='minisblack')
    else:
        path=tmp_path/'cells.png';Image.fromarray(np.repeat((source/150).astype(np.uint8)[...,None],3,axis=-1)).save(path)
    record=panel.import_image(path)
    if module=='fluoro':
        record['channels'][0]['name']='MYHC-A';record['channels'][1]['name']='MYHC-B'
    panel.update_measurement_scale_bar(record['uid'],um_per_px=.5,calibration_source='manual')
    window.set_module(module);controller.sync()
    return controller,panel,record,labels,path


def wait_myotube_job(window,controller):
    import time
    from PyQt6.QtTest import QTest
    deadline=time.monotonic()+15
    while window._analysis_job_service.is_busy(controller.job_key) and time.monotonic()<deadline:QTest.qWait(10)
    assert not window._analysis_job_service.is_busy(controller.job_key)


@pytest.mark.parametrize('module',['fluoro','brightfield'])
def test_independent_myotube_entry_results_review_and_exports(myotube_window,tmp_path,monkeypatch,module):
    from core.image_analysis import cellpose_backend
    from core.models import Project
    from openpyxl import load_workbook
    w=myotube_window
    assert hasattr(w,'myotube_controllers'), 'Independent muscle-tube workflows are missing'
    c,panel,record,labels,path=import_example(w,module,tmp_path)
    other=deepcopy(panel.analysis);calls=[]
    def worker(image,**settings):
        calls.append(image.copy());return labels,{'model':'cyto3'}
    monkeypatch.setattr(cellpose_backend,'run_cellpose_labels',worker)
    c.controls.entry.setChecked(True);c.controls.diameter.setValue(80);c.run();wait_myotube_job(w,c)
    result=c.current_result();assert result is not None,c.controls.summary.text()
    assert result['type']=='myotube_measurement' and result['modality']==module
    assert c.controls.table.rowCount()==1 and c.controls.table.columnCount()==8
    assert result['results'][0]['max_width_um']==pytest.approx(10,abs=1)
    assert result['params']['diameter_px']==80
    if module=='fluoro':
        w.fluoro_workspace.set_step(2);w.show();w.workflow_ui.sync()
        selector=w.fluoro_workspace.measure_task_combo
        selector.setCurrentIndex(selector.findData('myotube'));c.controls.entry.setChecked(True)
        from PyQt6.QtWidgets import QApplication
        QApplication.processEvents();c.show_overlay()
        kinds=[item.data(0) for item in w.fluoro_workspace.canvas._measurement_overlay_items]
        assert 'fluoro_myotube_width_overlay' in kinds,(c.controls.isVisible(),c.controls.entry.isChecked(),selector.currentData(),len(w.fluoro_workspace.canvas._image_items))
        assert 'fluoro_myotube_length_overlay' in kinds
    assert not panel.analysis.get('bright_round_cell_diameter_measurements')
    if module=='fluoro':assert calls[0].dtype==np.uint16 and calls[0].max()==25000
    for key,value in other.items():assert panel.analysis[key]==value
    original=deepcopy(result)
    dialog=c.open_review();dialog.session.replace_polygon(1,[[30,25],[220,25],[220,55],[30,55]])
    assert dialog.session.undo() and dialog.session.redo();dialog._apply()
    assert c.current_result()['results'][0]['max_width_um']>original['results'][0]['max_width_um']
    c.export_csv(tmp_path/f'{module}.csv');c.export_excel(tmp_path/f'{module}.xlsx');c.export_overlay(tmp_path/f'{module}.png')
    book=load_workbook(tmp_path/f'{module}.xlsx')
    assert book.active.cell(2,3).data_type=='n'
    assert book.active.cell(2,7).value==pytest.approx(c.current_result()['results'][0]['centerline_length_um'])
    book.close()
    import csv
    with (tmp_path/f'{module}.csv').open(encoding='utf-8-sig',newline='') as stream:
        exported=next(csv.DictReader(stream))
    assert float(exported['centerline_length_um'])==pytest.approx(c.current_result()['results'][0]['centerline_length_um'])
    archive=tmp_path/f'{module}.sfig';w.project.save_sfig(archive)
    restored=Project.load_sfig(archive)
    assert next(p for p in restored.panels if p.panel_type==panel.panel_type).analysis['myotube_measurements']==panel.analysis['myotube_measurements']
    restored.cleanup_archive_extract_root()


@pytest.mark.parametrize('module',['fluoro','brightfield'])
def test_myotube_image_channel_parameters_and_stale_jobs(myotube_window,tmp_path,monkeypatch,module):
    from core.image_analysis import cellpose_backend
    w=myotube_window
    assert hasattr(w,'myotube_controllers'), 'Independent muscle-tube workflows are missing'
    c,panel,first,labels,path=import_example(w,module,tmp_path)
    monkeypatch.setattr(cellpose_backend,'run_cellpose_labels',lambda *a,**k:(labels,dict(model='cyto3')))
    c.controls.diameter.setValue(90);c.run();wait_myotube_job(w,c)
    first_result=deepcopy(c.current_result())
    if module=='fluoro':
        c.controls.channel.setCurrentIndex(c.controls.channel.findData(1))
        assert c.current_result() is None and c.controls.table.rowCount()==0
        c.controls.diameter.setValue(35);c.run();wait_myotube_job(w,c)
        c.controls.channel.setCurrentIndex(c.controls.channel.findData(0))
        assert c.current_result()==first_result and c.controls.diameter.value()==90
    panel.import_image(path);c.sync()
    assert c.current_result() is None and c.controls.table.rowCount()==0
    panel.set_active_image(first['uid']);c.sync()
    assert c.current_result()==first_result and c.controls.diameter.value()==90
    pending=[];monkeypatch.setattr(w._analysis_job_service,'start',lambda **kwargs:pending.append(kwargs))
    c.run();assert pending[0]['is_current']()
    c.controls.diameter.setValue(120)
    assert not pending[0]['is_current']()
    assert c.current_result()==first_result


@pytest.mark.parametrize('module',['fluoro','brightfield'])
def test_myotube_footer_and_localization(myotube_window,module):
    w=myotube_window
    assert hasattr(w,'myotube_controllers'), 'Independent muscle-tube workflows are missing'
    w.set_module(module);c=w.myotube_controllers[module]
    if module=='fluoro':w.fluoro_workspace.set_step(2)
    else:w.data_panel.set_brightfield_step(2,emit_signal=False)
    c.controls.entry.setChecked(True);w.workflow_ui.sync()
    assert w.workflow_ui.footers[module].source is c.controls.run_button
    w.switch_language('zh');assert '肌管' in c.controls.entry.text()
    w.switch_language('en');assert 'myotube' in c.controls.entry.text().lower()

