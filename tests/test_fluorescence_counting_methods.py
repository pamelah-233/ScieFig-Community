import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import json
import numpy as np
import pytest
from openpyxl import load_workbook
from PyQt6.QtWidgets import QApplication
from core.models import FluorescencePanel


def make_panel():
    yy,xx=np.indices((120,180))
    image=sum(150*np.exp(-((xx-x)**2+(yy-55)**2)/(2*6**2)) for x in (40,110)).astype(np.float32)
    panel=FluorescencePanel(panel_id='method-test')
    panel.data['fluorescence_images']=[dict(uid='x',source_name='test.tif',channels=[dict(index=0,name='Signal')])]
    panel.data['active_image_uid']='x'
    panel.black_balanced_image_data=lambda _:image[None]
    return panel


def test_log_count_has_parameters_but_no_segmentation_area():
    panel=make_panel()
    row=panel.count_channel_objects('x',[0],algorithm='log',diameter_px=20,sensitivity=1)[0]
    assert row['count']==2
    assert row['algorithm']=='log'
    assert row['algorithm_parameters']['diameter_px']==20
    assert row['total_positive_area_px'] is None
    assert row['mean_object_area_px'] is None
    json.dumps(row)


def test_selector_and_excel_use_saved_method_not_current_selection(tmp_path):
    from ui.panels.fluorescence_workspace import FluorescenceWorkspace
    from ui.widgets.measurement_excel_export import export_table_xlsx
    app=QApplication.instance() or QApplication([])
    panel=make_panel()
    panel.count_channel_objects('x',[0],algorithm='log',diameter_px=20)
    workspace=FluorescenceWorkspace()
    try:
        workspace.panel=panel
        assert workspace.count_algorithm_combo.currentData()=='traditional'
        workspace.count_algorithm_combo.setCurrentIndex(workspace.count_algorithm_combo.findData('cellpose_cyto3'))
        assert not workspace.sample_spot_button.isEnabled()
        workspace._populate_count_results()
        assert workspace.count_table.item(0,5).text()=='LoG'
        target=export_table_xlsx(workspace.count_table,tmp_path/'methods.xlsx',(1,2,3,4))
        book=load_workbook(target);ws=book.active
        assert ws.cell(2,6).value=='LoG'
        assert json.loads(ws.cell(2,7).value)['diameter_px']==20
        assert ws.cell(2,8).value=='test.tif'
        assert ws.cell(2,3).value is None
        book.close()
    finally:workspace.close()


def test_statistics_do_not_mix_methods():
    from core.fluoro_statistics import collect_fluoro_stat_values
    panel=make_panel()
    panel.count_channel_objects('x',[0],algorithm='traditional')
    panel.count_channel_objects('x',[0],algorithm='log',diameter_px=20)
    rows=collect_fluoro_stat_values(panel,source='object_count',metric='count')
    assert len(rows)==2
    assert len({r['sample_key'] for r in rows})==2
    assert all(r['n_measurements']==1 for r in rows)


def test_each_count_refreshes_and_selects_new_numeric_row():
    from ui.panels.fluorescence_workspace import FluorescenceWorkspace
    app=QApplication.instance() or QApplication([])
    panel=make_panel();workspace=FluorescenceWorkspace()
    try:
        workspace.panel=panel
        workspace._effective_count_channels=lambda:[0]
        workspace._populate_overlap_results=lambda:None
        workspace._populate_fluoro_stats_controls=lambda:None
        workspace.threshold_method_combo.setCurrentText('manual')
        workspace.threshold_value_spin.setValue(20)
        workspace._count_spots()
        assert workspace.count_table.item(0,1).text()=='2'
        workspace.threshold_value_spin.setValue(200)
        workspace._count_spots()
        assert workspace.count_table.rowCount()==2
        assert workspace.count_table.item(1,1).text()=='0'
        assert workspace.count_table.currentRow()==1
    finally:workspace.close()


def test_new_counts_clear_old_statistical_numbers():
    from ui.panels.fluorescence_workspace import FluorescenceWorkspace
    app=QApplication.instance() or QApplication([])
    panel=make_panel();workspace=FluorescenceWorkspace()
    try:
        workspace.panel=panel
        panel.analysis['fluoro_statistics']={'groups':[dict(group='old',n=4,mean=123,sem=1)]}
        workspace._populate_fluoro_stats_controls()
        assert workspace.fluoro_stats_result_table.rowCount()==1
        panel.count_channel_objects('x',[0])
        workspace._populate_fluoro_stats_controls()
        assert workspace.fluoro_stats_result_table.rowCount()==0
    finally:workspace.close()


def test_different_parameters_do_not_get_same_default_statistics_group():
    from core.fluoro_statistics import collect_fluoro_stat_values,compute_fluoro_statistics
    panel=make_panel()
    panel.count_channel_objects('x',[0],algorithm='log',diameter_px=20)
    panel.count_channel_objects('x',[0],algorithm='log',diameter_px=40)
    rows=collect_fluoro_stat_values(panel,'object_count','count')
    assignments={r['sample_group_key']:r['lane_label'] for r in rows}
    result=compute_fluoro_statistics(panel,'object_count','count',assignments,rows[0]['lane_label'],'Unpaired t-test')
    assert len(result['groups'])==2
    assert all(group['n']==1 for group in result['groups'])
