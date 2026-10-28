from types import SimpleNamespace
from pathlib import Path


def test_provenance_keeps_method_calibration_review_and_pair_ids():
    from core.measurement_provenance import measurement_provenance_rows
    panel = SimpleNamespace(analysis={'object_counts': [dict(count_uuid='c1',image_uid='i1',algorithm='log',
        algorithm_parameters={'sensitivity':1.2},measurement_scale_bar={'um_per_px':.5},object_review={'revision':1})],
        'fluoro_statistics': {'plot_data': [dict(sample_key='i1:0',pair_id='mouse1',group='A',value=8)]}},
        images=lambda: [dict(uid='i1',source_name='sample',source_path='C:/synthetic/sample.tif')])
    rows = measurement_provenance_rows(panel)
    assert rows[0]['measurement_id'] == 'c1'
    assert rows[0]['method'] == 'log'
    assert rows[0]['calibration'] == {'um_per_px':.5}
    assert rows[0]['review'] == {'revision':1}
    assert rows[-1]['pair_id'] == 'mouse1'


def test_workbook_exports_provenance_without_formulas(tmp_path):
    import os
    os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
    from PyQt6.QtWidgets import QApplication, QTableWidget, QTableWidgetItem
    from openpyxl import load_workbook
    from ui.widgets.measurement_excel_export import export_tables_xlsx
    app = QApplication.instance() or QApplication([])
    table = QTableWidget(1,1)
    table.setHorizontalHeaderLabels(['Count'])
    table.setItem(0,0,QTableWidgetItem('3'))
    path = export_tables_xlsx([('Counts',table,(0,))], tmp_path/'trace.xlsx', provenance=[
        dict(image_name='=HYPERLINK("bad")',method='log',parameters={'sensitivity':1.2},pair_id='01')])
    book = load_workbook(path)
    sheet = book['Provenance']
    assert sheet['A2'].data_type == 's'
    assert sheet['D2'].value == '01'
    assert '1.2' in sheet['C2'].value
    book.close()


def test_ai_ids_and_stored_settings_are_exported():
    from core.measurement_provenance import measurement_provenance_rows
    panel = SimpleNamespace(images=lambda: [], analysis={'fluorescence_ai_assisted_measurements': [
        dict(ai_measurement_uuid='ai-1', metadata={'model':'cyto3'}, batch_provenance={'settings':{'sensitivity':2}})]})
    row = measurement_provenance_rows(panel)[0]
    assert row['measurement_id'] == 'ai-1'
    assert row['parameters'] == {'sensitivity':2}
    assert row['metadata']['model'] == 'cyto3'


def test_long_review_provenance_is_losslessly_chunked(tmp_path):
    import os, json
    os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
    from PyQt6.QtWidgets import QApplication, QTableWidget, QTableWidgetItem
    from openpyxl import load_workbook
    from ui.widgets.measurement_excel_export import export_tables_xlsx
    app=QApplication.instance() or QApplication([])
    table=QTableWidget(1,1)
    table.setItem(0,0,QTableWidgetItem('1'))
    review={'audit':'记录' * 24000}
    path=export_tables_xlsx([('M',table,(0,))],tmp_path/'long.xlsx',provenance=[{'review':review}])
    book=load_workbook(path)
    parts=[r[3] for r in book['Provenance values'].iter_rows(min_row=2,values_only=True)]
    assert json.loads(''.join(parts)) == review
    book.close()
